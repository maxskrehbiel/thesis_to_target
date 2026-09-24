"""CSV and Excel writers for targets, universe, evidence ledger, review queue and raw records."""

from __future__ import annotations

import csv
import re
import zipfile
from collections.abc import Mapping, Sequence
from dataclasses import asdict, fields
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .evidence import render_value
from .formatting import usd_millions, whole_thousands
from .models import Company
from .pipeline import RunResult

_MAX_COLUMN_WIDTH = 60
_MIN_COLUMN_WIDTH = 10
# Rows inspected when sizing a column; later rows rarely change the width.
_WIDTH_SAMPLE_ROWS = 200
# A fixed document timestamp keeps the workbook byte-identical from run to run.
_DOC_TIMESTAMP = datetime(2026, 1, 1)
_MODIFIED_RE = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*")

COMPANY_COLUMNS = (
    "data",
    "rank",
    "tier",
    "company_id",
    "name",
    "legal_name",
    "city",
    "state",
    "fit_tier",
    "segment",
    "service_mix",
    "employees_range",
    "ebitda_median",
    "ebitda_95_low",
    "ebitda_95_high",
    "p_ebitda_ge_min",
    "p_ebitda_in_band",
    "size_confidence",
    "obviousness",
    "non_obvious",
    "n_sources",
    "sources",
    "website",
    "composite",
    "exclusions",
    "flags",
    "ledger_claims",
)
LEDGER_COLUMNS = (
    "claim_id",
    "company_id",
    "company",
    "claim",
    "value",
    "evidence",
    "source",
    "confidence",
    "kind",
    "method",
    "record_ids",
    "depends_on",
    "external_ref",
)
REVIEW_COLUMNS = (
    "record_a",
    "record_b",
    "name_a",
    "name_b",
    "source_a",
    "source_b",
    "state_a",
    "state_b",
    "score",
    "name_score",
    "reasons",
    "decision",
)
RAW_COLUMNS = (
    "data",
    "record_id",
    "source",
    "name",
    "dba",
    "city",
    "state",
    "zip",
    "phone",
    "website",
    "attributes",
)


def data_label(result: RunResult) -> str:
    """``"synthetic"`` when every source served generated data, else ``"source"``."""
    return "synthetic" if result.synthetic else "source"


def company_row(company: Company, claim_span: str, label: str) -> dict[str, Any]:
    """Flatten one company into an export row.

    Args:
        company: Company with stage results attached.
        claim_span: First-to-last ledger claim ids about the company.
        label: Value of the ``data`` column (see :func:`data_label`).

    Returns:
        Column name -> value, covering :data:`COMPANY_COLUMNS`.
    """
    est, fit, presence = company.estimate, company.fit, company.presence
    return {
        "data": label,
        "rank": company.rank,
        "tier": company.tier,
        "company_id": company.company_id,
        "name": company.name,
        "legal_name": company.legal_name,
        "city": company.city,
        "state": company.state,
        "fit_tier": fit.tier if fit else None,
        "segment": fit.segment if fit else None,
        "service_mix": fit.service_mix if fit else None,
        "employees_range": f"{est.employees_lo:.0f}-{est.employees_hi:.0f}" if est else None,
        "ebitda_median": whole_thousands(est.ebitda_mid) if est else None,
        "ebitda_95_low": whole_thousands(est.ebitda_lo) if est else None,
        "ebitda_95_high": whole_thousands(est.ebitda_hi) if est else None,
        "p_ebitda_ge_min": round(est.p_ge_min, 3) if est else None,
        "p_ebitda_in_band": round(est.p_in_band, 3) if est else None,
        "size_confidence": est.confidence if est else "unsized",
        "obviousness": round(presence.obviousness, 3) if presence else None,
        "non_obvious": presence.non_obvious if presence else None,
        "n_sources": len(company.sources),
        "sources": "; ".join(company.sources),
        "website": company.domain,
        "composite": company.composite,
        "exclusions": "; ".join(company.exclusions),
        "flags": "; ".join(company.flags),
        "ledger_claims": claim_span,
    }


def _claim_span(result: RunResult, company_id: str) -> str:
    claims = result.ledger.for_company(company_id)
    return f"{claims[0].claim_id}..{claims[-1].claim_id}" if claims else ""


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    return value


def write_csv(path: Path, columns: Sequence[str], rows: Sequence[Mapping[str, Any]]) -> Path:
    """Write rows to a UTF-8 CSV with a fixed column order.

    Args:
        path: Output file.
        columns: Column order.
        rows: Row mappings; missing keys become empty cells.

    Returns:
        ``path``.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(
            fh, fieldnames=list(columns), extrasaction="ignore", lineterminator="\n"
        )
        writer.writeheader()
        for row in rows:
            writer.writerow({k: _cell(row.get(k)) for k in columns})
    return path


def _sheet(
    wb: Workbook, title: str, columns: Sequence[str], rows: Sequence[Mapping[str, Any]]
) -> None:
    ws = wb.create_sheet(title)
    ws.append(list(columns))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([_cell(row.get(c)) for c in columns])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    for idx, column in enumerate(columns, start=1):
        longest = max(
            [len(str(column)), *(len(str(_cell(r.get(column)))) for r in rows[:_WIDTH_SAMPLE_ROWS])]
        )
        ws.column_dimensions[get_column_letter(idx)].width = min(
            _MAX_COLUMN_WIDTH, max(_MIN_COLUMN_WIDTH, longest + 2)
        )


def _assumption_rows(result: RunResult) -> list[dict[str, Any]]:
    t = result.thesis
    rows: list[dict[str, Any]] = [
        {
            "parameter": "EBITDA band",
            "value": f"{usd_millions(t.ebitda_min)} to {usd_millions(t.ebitda_max)}",
        },
        {"parameter": "states", "value": ", ".join(t.states)},
        {"parameter": "NAICS codes", "value": ", ".join(t.naics)},
    ]
    for seg in t.segments.values():
        lo, hi = seg.revenue_per_employee
        rows.append(
            {"parameter": f"revenue per employee: {seg.name}", "value": f"${lo:,.0f} to ${hi:,.0f}"}
        )
    for mix, (lo, hi) in t.margins.items():
        rows.append({"parameter": f"EBITDA margin: {mix}", "value": f"{lo:.0%} to {hi:.0%}"})
    for spec in t.presence_databases:
        rows.append({"parameter": f"obviousness weight: {spec.database}", "value": spec.weight})
    for section in (t.resolve, t.estimation, t.presence, t.ranking):
        for f in fields(section):
            value = getattr(section, f.name)
            if f.name == "generic_tokens":
                value = ", ".join(sorted(value))
            rows.append(
                {"parameter": f"{type(section).__name__}.{f.name}", "value": render_value(value)}
            )
    return rows


def _qa_rows(result: RunResult) -> list[dict[str, Any]]:
    if result.qa is None:
        return [
            {"metric": "status", "value": "not available (the sources supplied no ground truth)"}
        ]
    rows = [
        {"metric": f"resolution.{k}", "value": v} for k, v in asdict(result.qa.resolution).items()
    ]
    rows += [{"metric": k, "value": v} for k, v in asdict(result.qa).items() if k != "resolution"]
    return rows


def _about_rows(result: RunResult) -> list[dict[str, Any]]:
    t, f = result.thesis, result.funnel
    if result.synthetic:
        data = (
            f"synthetic: every company, record and listing is fictional (seed {t.synthetic.seed})"
        )
    else:
        data = "records supplied by the configured source adapters"
    rows: list[dict[str, Any]] = [
        {"item": "data", "value": data},
        {"item": "thesis", "value": t.name},
        {"item": "description", "value": t.description},
    ]
    rows += [{"item": f"funnel: {k}", "value": render_value(v)} for k, v in asdict(f).items()]
    return rows


def review_rows(result: RunResult) -> list[dict[str, Any]]:
    """Rows of the human review queue, with blank decisions to fill in."""
    by_id = {r.record_id: r for r in result.records}
    return [
        {
            "record_a": p.a,
            "record_b": p.b,
            "name_a": by_id[p.a].name,
            "name_b": by_id[p.b].name,
            "source_a": by_id[p.a].source,
            "source_b": by_id[p.b].source,
            "state_a": by_id[p.a].state,
            "state_b": by_id[p.b].state,
            "score": p.score,
            "name_score": p.name_score,
            "reasons": "; ".join(p.reasons),
            "decision": "",
        }
        for p in result.review_queue
    ]


def freeze_workbook(path: Path) -> None:
    """Rewrite an .xlsx with fixed timestamps so the same content gives the same bytes.

    openpyxl stamps the save time into every zip entry and into ``dcterms:modified``.

    Args:
        path: Workbook to rewrite in place.
    """
    with zipfile.ZipFile(path) as src:
        entries = [(info.filename, src.read(info.filename)) for info in src.infolist()]
    stamp = _DOC_TIMESTAMP.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    date_time = (_DOC_TIMESTAMP.year, _DOC_TIMESTAMP.month, _DOC_TIMESTAMP.day, 0, 0, 0)
    with zipfile.ZipFile(path, "w") as dst:
        for name, data in entries:
            if name == "docProps/core.xml":
                data = _MODIFIED_RE.sub(rb"\g<1>" + stamp, data)
            info = zipfile.ZipInfo(name, date_time=date_time)
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, data)


def export_run(result: RunResult, out_dir: str | Path) -> dict[str, Path]:
    """Write every output file of a run.

    Args:
        result: A finished pipeline run.
        out_dir: Directory to write into (created if missing).

    Returns:
        Output name -> path written.
    """
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    names = {c.company_id: c.name for c in result.companies}
    label = data_label(result)
    ranked = [company_row(c, _claim_span(result, c.company_id), label) for c in result.ranked()]
    universe = [company_row(c, _claim_span(result, c.company_id), label) for c in result.companies]
    ledger = result.ledger.rows(names)
    review = review_rows(result)
    raw = [
        {
            "data": label,
            **{k: v for k, v in asdict(r).items() if k != "attributes"},
            "attributes": render_value(r.attributes),
        }
        for r in result.records
    ]

    paths = {
        "targets_csv": write_csv(out / "targets.csv", COMPANY_COLUMNS, ranked),
        "evidence_ledger_csv": write_csv(out / "evidence_ledger.csv", LEDGER_COLUMNS, ledger),
        "review_queue_csv": write_csv(out / "review_queue.csv", REVIEW_COLUMNS, review),
        "raw_records_csv": write_csv(out / "raw_records.csv", RAW_COLUMNS, raw),
    }

    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "about", ("item", "value"), _about_rows(result))
    _sheet(wb, "targets", COMPANY_COLUMNS, ranked)
    _sheet(wb, "universe", COMPANY_COLUMNS, universe)
    _sheet(wb, "evidence_ledger", LEDGER_COLUMNS, ledger)
    _sheet(wb, "review_queue", REVIEW_COLUMNS, review)
    _sheet(wb, "raw_records", RAW_COLUMNS, raw)
    _sheet(wb, "assumptions", ("parameter", "value"), _assumption_rows(result))
    _sheet(wb, "qa", ("metric", "value"), _qa_rows(result))
    wb.properties.creator = "thesis_to_target"
    wb.properties.created = _DOC_TIMESTAMP
    xlsx = out / "targets.xlsx"
    wb.save(xlsx)
    freeze_workbook(xlsx)
    paths["targets_xlsx"] = xlsx
    return paths
