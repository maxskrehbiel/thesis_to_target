"""CSV and Excel exports."""

from __future__ import annotations

import csv
import dataclasses
import io
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from thesis_to_target.export import COMPANY_COLUMNS, LEDGER_COLUMNS, data_label, export_run
from thesis_to_target.pipeline import RunResult


def _rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def test_export_writes_every_file(tmp_path: Path, demo_run: RunResult) -> None:
    paths = export_run(demo_run, tmp_path / "out")
    assert {p.name for p in paths.values()} == {
        "targets.csv",
        "evidence_ledger.csv",
        "review_queue.csv",
        "raw_records.csv",
        "targets.xlsx",
    }
    targets = _rows(paths["targets_csv"])
    assert list(targets[0]) == list(COMPANY_COLUMNS)
    assert len(targets) == demo_run.funnel.in_scope
    assert all(row["data"] == "synthetic" for row in targets)
    assert [int(row["rank"]) for row in targets] == list(range(1, len(targets) + 1))
    ledger = _rows(paths["evidence_ledger_csv"])
    assert list(ledger[0]) == list(LEDGER_COLUMNS) and len(ledger) == len(demo_run.ledger)
    raw = _rows(paths["raw_records_csv"])
    assert len(raw) == len(demo_run.records) and all(r["data"] == "synthetic" for r in raw)
    review = _rows(paths["review_queue_csv"])
    assert len(review) == len(demo_run.review_queue) and all(r["decision"] == "" for r in review)


def test_every_target_cell_traces_to_the_ledger(tmp_path: Path, demo_run: RunResult) -> None:
    targets = _rows(export_run(demo_run, tmp_path)["targets_csv"])
    claim_ids = {c.claim_id for c in demo_run.ledger}
    for row in targets:
        first, last = row["ledger_claims"].split("..")
        assert first in claim_ids and last in claim_ids and first.split(".")[0] == row["company_id"]


def test_workbook_sheets_and_metadata(tmp_path: Path, demo_run: RunResult) -> None:
    xlsx = export_run(demo_run, tmp_path)["targets_xlsx"]
    wb = load_workbook(xlsx, read_only=True)
    assert wb.sheetnames == [
        "about",
        "targets",
        "universe",
        "evidence_ledger",
        "review_queue",
        "raw_records",
        "assumptions",
        "qa",
    ]
    assert wb.properties.creator == "thesis_to_target"
    about = [row for row in wb["about"].iter_rows(values_only=True)]
    assert any(str(value).startswith("synthetic: every company") for row in about for value in row)


def test_workbook_bytes_are_reproducible(tmp_path: Path, demo_run: RunResult) -> None:
    first = export_run(demo_run, tmp_path / "a")["targets_xlsx"].read_bytes()
    second = export_run(demo_run, tmp_path / "b")["targets_xlsx"].read_bytes()
    assert first == second
    with zipfile.ZipFile(io.BytesIO(first)) as archive:
        core = archive.read("docProps/core.xml")
        assert b"<dcterms:modified" in core and b"2026-01-01T00:00:00Z</dcterms:modified>" in core
        assert {info.date_time for info in archive.infolist()} == {(2026, 1, 1, 0, 0, 0)}


def test_data_label_follows_the_sources(tmp_path: Path, demo_run: RunResult) -> None:
    assert data_label(demo_run) == "synthetic"
    sourced = dataclasses.replace(demo_run, synthetic=False, qa=None)
    assert data_label(sourced) == "source"
    paths = export_run(sourced, tmp_path)
    assert {row["data"] for row in _rows(paths["targets_csv"])} == {"source"}
    wb = load_workbook(paths["targets_xlsx"], read_only=True)
    cells = [str(v) for row in wb["about"].iter_rows(values_only=True) for v in row]
    assert not any("synthetic" in c.lower() for c in cells)
    qa_cells = [str(v) for row in wb["qa"].iter_rows(values_only=True) for v in row]
    assert any("no ground truth" in c for c in qa_cells)


def test_money_columns_are_whole_thousands(tmp_path: Path, demo_run: RunResult) -> None:
    for row in _rows(export_run(demo_run, tmp_path)["targets_csv"]):
        if row["ebitda_median"]:
            assert int(row["ebitda_95_low"]) % 1000 == 0 and int(row["ebitda_median"]) % 1000 == 0
